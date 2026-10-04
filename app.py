import io
import streamlit as st
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from docx import Document

st.set_page_config(
    page_title="Asistente IA · Proyectos Viales",
    page_icon="🛣️",
    layout="wide",
    initial_sidebar_state="expanded",
)

STATES = [
    "Cumple",
    "Cumple parcialmente",
    "No cumple",
    "No se encontró información",
    "Requiere revisión especializada",
]

DEMO_DOCS = [
    {"Documento": "01_Memoria_Descriptiva_Demo.pdf", "Componente": "Memoria descriptiva", "Estado": "Reconocido"},
    {"Documento": "02_Metrados_Demo.xlsx", "Componente": "Metrados", "Estado": "Reconocido"},
    {"Documento": "03_Presupuesto_Demo.xlsx", "Componente": "Presupuesto", "Estado": "Reconocido"},
]

DEMO_REQS = [
    {
        "id": "REQ-01", "componente": "Memoria descriptiva",
        "requisito": "El entregable identifica el proyecto y describe su alcance.",
        "criterio": "Lista de chequeo del caso de prueba",
        "evidencia": "La memoria identifica el proyecto vial y describe el alcance general de la intervención.",
        "fuente": "01_Memoria_Descriptiva_Demo.pdf · sección 1",
        "estado": "Cumple",
        "obs": "Se identifica el proyecto y su alcance en la documentación revisada.",
        "peso": 25,
    },
    {
        "id": "REQ-02", "componente": "Metrados",
        "requisito": "El componente de metrados forma parte de la documentación presentada.",
        "criterio": "Lista de chequeo del caso de prueba",
        "evidencia": "Se encuentra el archivo 02_Metrados_Demo.xlsx dentro del conjunto documental.",
        "fuente": "02_Metrados_Demo.xlsx · hoja Metrados",
        "estado": "Cumple",
        "obs": "Se encontró el componente de metrados dentro de la documentación de prueba.",
        "peso": 20,
    },
    {
        "id": "REQ-03", "componente": "Presupuesto",
        "requisito": "El presupuesto presenta información relacionada con las partidas del caso de prueba.",
        "criterio": "Lista de chequeo del caso de prueba",
        "evidencia": "El documento contiene partidas y montos de demostración asociados al proyecto.",
        "fuente": "03_Presupuesto_Demo.xlsx · hoja Presupuesto",
        "estado": "Cumple parcialmente",
        "obs": "Se encontró información de presupuesto; corresponde al especialista validar su consistencia técnica.",
        "peso": 20,
    },
    {
        "id": "REQ-04", "componente": "Cronograma",
        "requisito": "Se presenta el cronograma correspondiente al entregable revisado.",
        "criterio": "Lista de chequeo del caso de prueba",
        "evidencia": "No se localizó un archivo de cronograma dentro del conjunto de demostración.",
        "fuente": "Conjunto documental cargado",
        "estado": "No se encontró información",
        "obs": "No se encontró evidencia del cronograma en los documentos de prueba. No se interpreta automáticamente como incumplimiento.",
        "peso": 15,
    },
    {
        "id": "REQ-05", "componente": "Planos",
        "requisito": "La revisión del plano permite verificar aspectos que requieren interpretación técnica visual.",
        "criterio": "Alcance del MVP",
        "evidencia": "La verificación requiere interpretación especializada de planos, función excluida del MVP.",
        "fuente": "Alcance definido para el MVP",
        "estado": "Requiere revisión especializada",
        "obs": "Derivar al especialista para revisión técnica del plano.",
        "peso": 20,
    },
]

# ---------- Estado ----------
if "step" not in st.session_state:
    st.session_state.step = 1
if "results" not in st.session_state:
    st.session_state.results = [dict(x, validacion="Pendiente") for x in DEMO_REQS]
if "project" not in st.session_state:
    st.session_state.project = "Proyecto Vial Demostrativo"
if "deliverable" not in st.session_state:
    st.session_state.deliverable = "Expediente técnico · caso de prueba"
if "mode" not in st.session_state:
    st.session_state.mode = "Usar documentos de prueba incluidos"

# ---------- Estilo ----------
st.markdown("""
<style>
:root{
  --azul:#124f8c; --azul2:#1d6dac; --azul-suave:#eaf3fc;
  --verde:#23864b; --gris:#5c6b7a; --borde:#dbe4ee; --fondo:#f7f9fc;
}
.block-container{padding-top:1.0rem;max-width:1420px;padding-bottom:2rem}
[data-testid="stSidebar"]{background:#f3f6fa}
.hero{
  background:linear-gradient(110deg,#0d427c,#216fae);
  padding:24px 30px;border-radius:18px;color:white;margin-bottom:16px;
  box-shadow:0 4px 14px rgba(13,66,124,.12)
}
.hero h1{margin:0;font-size:2.15rem;letter-spacing:.2px}
.hero p{margin:.35rem 0 0;font-size:1.02rem}
.badge{
  display:inline-block;background:#e9f3ff;color:#0d4f8c;border:1px solid #bfd8f1;
  border-radius:999px;padding:5px 10px;font-size:.82rem;font-weight:700;margin-top:10px
}
.step{
  padding:10px 8px;border-radius:10px;background:#f0f4f8;text-align:center;
  font-weight:700;color:#27384a;border:1px solid transparent
}
.active{background:#dcecff;color:#0b3b75;border-color:#78afe5}
.section-card{
  padding:18px 20px;border:1px solid var(--borde);border-radius:14px;background:white;
  box-shadow:0 2px 8px rgba(30,55,80,.04);margin-bottom:12px
}
.trace{
  background:#f8fbff;border-left:4px solid #2a73b7;padding:12px 14px;border-radius:8px;
  margin:8px 0 12px
}
.small{font-size:.88rem;color:#5c6b7a}
div[data-testid="stMetric"]{
  background:white;border:1px solid var(--borde);padding:12px;border-radius:12px
}
.stButton>button[kind="primary"]{background:#1769aa}
hr{border:none;border-top:1px solid #e7edf3}
</style>
""", unsafe_allow_html=True)

st.markdown("""
<div class="hero">
  <h1>ASISTENTE IA</h1>
  <p>Revisión y Evaluación de Perfiles y Expedientes Técnicos de Proyectos Viales</p>
  <span class="badge">MVP demostrativo · Hito 2 · V2.1</span>
</div>
""", unsafe_allow_html=True)

labels = ["1 · Inicio", "2 · Carga", "3 · Procesamiento", "4 · Revisión", "5 · Resultados"]
cols = st.columns(5)
for i, (c, label) in enumerate(zip(cols, labels), 1):
    c.markdown(
        f'<div class="step {"active" if st.session_state.step == i else ""}">{label}</div>',
        unsafe_allow_html=True,
    )
st.caption("Flujo: documentos → requisitos → evidencia → análisis preliminar → validación humana → productos de salida.")

# ---------- Utilidades ----------
def nav(back=None, next_=None, next_label="Continuar"):
    a, b = st.columns([1, 1])
    if back and a.button("← Anterior", use_container_width=True):
        st.session_state.step = back
        st.rerun()
    if next_ and b.button(next_label + " →", type="primary", use_container_width=True):
        st.session_state.step = next_
        st.rerun()

def progress_pct(rows):
    score = {
        "Cumple": 1,
        "Cumple parcialmente": 0.5,
        "No cumple": 0,
        "No se encontró información": 0,
        "Requiere revisión especializada": 0,
    }
    den = sum(r["peso"] for r in rows)
    return 100 * sum(r["peso"] * score.get(r["estado"], 0) for r in rows) / den if den else 0

def excel_bytes(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Matriz de revisión"
    headers = [
        "ID","Componente","Requisito","Criterio","Evidencia","Fuente",
        "Resultado / Estado","Observación","Peso (%)","Validación humana"
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="124F8C")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append([
            r["id"],r["componente"],r["requisito"],r["criterio"],r["evidencia"],
            r["fuente"],r["estado"],r["obs"],r["peso"],r.get("validacion","Pendiente")
        ])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [12,22,45,32,48,34,30,48,12,20]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64+i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()

def word_bytes(rows, pct):
    d = Document()
    d.add_heading("Borrador de informe de revisión", 0)
    d.add_paragraph(f"Proyecto: {st.session_state.project}")
    d.add_paragraph(f"Entregable: {st.session_state.deliverable}")
    d.add_heading("Resumen", 1)
    d.add_paragraph(
        f"Porcentaje de avance demostrativo: {pct:.1f} %. "
        "Los resultados son preliminares y requieren validación del especialista."
    )
    d.add_heading("Resultados de revisión", 1)
    for r in rows:
        d.add_heading(f"{r['id']} · {r['componente']}", 2)
        d.add_paragraph("Requisito: " + r["requisito"])
        d.add_paragraph("Criterio: " + r["criterio"])
        d.add_paragraph("Evidencia: " + r["evidencia"])
        d.add_paragraph("Fuente: " + r["fuente"])
        d.add_paragraph("Estado: " + r["estado"])
        d.add_paragraph("Observación: " + r["obs"])
        d.add_paragraph("Validación humana: " + r.get("validacion", "Pendiente"))
    bio = io.BytesIO()
    d.save(bio)
    return bio.getvalue()

# ---------- Pantallas ----------
if st.session_state.step == 1:
    st.header("1. Inicio / Nueva revisión")
    st.info("Este prototipo utiliza un caso de prueba delimitado. No emite conformidad ni sustituye el criterio técnico del especialista.")
    c1, c2 = st.columns(2)
    st.session_state.project = c1.text_input("Proyecto / expediente", st.session_state.project)
    st.session_state.deliverable = c2.text_input("Entregable / revisión", st.session_state.deliverable)

    st.markdown("""
    <div class="trace"><b>Recorrido disponible</b><br>
    Carga → clasificación → procesamiento → evidencia → resultado preliminar → validación humana → Excel y Word.
    </div>
    """, unsafe_allow_html=True)

    if st.button("＋ NUEVA REVISIÓN", type="primary", use_container_width=True):
        st.session_state.results = [dict(x, validacion="Pendiente") for x in DEMO_REQS]
        st.session_state.step = 2
        st.rerun()

elif st.session_state.step == 2:
    st.header("2. Carga y clasificación de documentos")
    st.session_state.mode = st.radio(
        "Modo de demostración",
        ["Usar documentos de prueba incluidos", "Cargar mis propios archivos"],
        horizontal=True,
        index=0 if st.session_state.mode.startswith("Usar") else 1,
    )

    if st.session_state.mode.startswith("Usar"):
        st.success("Caso de prueba preparado. Los documentos son simulados y sirven únicamente para demostrar el flujo del MVP.")
        st.dataframe(
            pd.DataFrame(DEMO_DOCS),
            use_container_width=True,
            hide_index=True,
            column_config={
                "Documento": st.column_config.TextColumn("Documento", width="large"),
                "Componente": st.column_config.TextColumn("Componente", width="medium"),
                "Estado": st.column_config.TextColumn("Estado", width="small"),
            },
        )
    else:
        files = st.file_uploader(
            "PDF con texto reconocible, Word o Excel",
            type=["pdf", "docx", "xlsx"],
            accept_multiple_files=True,
        )
        if files:
            st.success(f"{len(files)} archivo(s) recibido(s).")
            uploaded = pd.DataFrame([
                {"Documento": f.name, "Componente": "Por confirmar", "Estado": "Recibido"}
                for f in files
            ])
            st.dataframe(uploaded, use_container_width=True, hide_index=True)
            st.warning("Para mantener la demostración estable, el análisis posterior continúa utilizando la matriz de prueba predefinida.")

    st.markdown("#### Criterios de revisión")
    c1, c2, c3 = st.columns(3)
    c1.selectbox("Lista de chequeo", ["Lista de chequeo · caso de prueba"])
    c2.selectbox("Requisitos", ["Requisitos · caso de prueba"])
    c3.selectbox("Normativa / criterios", ["Criterios incorporados · demostración"])

    st.caption("Antes de iniciar el análisis, el especialista confirma los documentos y criterios que serán utilizados.")
    nav(1, 3, "Iniciar revisión")

elif st.session_state.step == 3:
    st.header("3. Procesamiento")
    stages = [
        "Recepción y validación de archivos",
        "Identificación y clasificación documental",
        "Extracción y estructuración",
        "Búsqueda de evidencia por requisito",
        "Análisis preliminar asistido por IA",
        "Consolidación de resultados",
    ]
    for i, name in enumerate(stages, 1):
        st.markdown(f"✅ **{i}. {name}** — Completado")
    st.progress(1.0)
    st.success("Procesamiento demostrativo finalizado. Los resultados están listos para revisión humana.")
    st.info("Modo demostrativo: los resultados son predefinidos para que el recorrido sea estable y reproducible. La arquitectura prevista incorpora Gemini API para el análisis asistido.")
    nav(2, 4, "Revisar resultados")

elif st.session_state.step == 4:
    st.header("4. Revisión y validación de resultados")
    rows = st.session_state.results

    pendientes = sum(r.get("validacion", "Pendiente") == "Pendiente" for r in rows)
    m1, m2, m3 = st.columns(3)
    m1.metric("Requisitos analizados", len(rows))
    m2.metric("Pendientes de validación humana", pendientes)
    m3.metric("Avance demostrativo", f"{progress_pct(rows):.1f} %")

    ids = [f"{r['id']} · {r['componente']}" for r in rows]
    selected = st.selectbox("Seleccione un requisito", ids)
    idx = ids.index(selected)
    r = rows[idx]

    left, right = st.columns([1.2, 1])
    with left:
        st.markdown("#### Requisito")
        st.write(r["requisito"])
        st.markdown("#### Criterio / norma")
        st.write(r["criterio"])
        st.markdown("#### Evidencia encontrada")
        st.info(r["evidencia"])
        st.markdown(f"**Fuente:** {r['fuente']}")

    with right:
        st.markdown("#### Resultado preliminar")
        new_state = st.selectbox(
            "Estado",
            STATES,
            index=STATES.index(r["estado"]),
            key="state_" + r["id"],
        )
        new_obs = st.text_area(
            "Observación preliminar",
            r["obs"],
            height=155,
            key="obs_" + r["id"],
        )
        st.caption(f"Validación actual: **{r.get('validacion','Pendiente')}**")

        c1, c2, c3 = st.columns(3)
        if c1.button("✓ Aceptar", key="acc_" + r["id"], use_container_width=True):
            rows[idx]["estado"] = new_state
            rows[idx]["obs"] = new_obs
            rows[idx]["validacion"] = "Aceptado"
            st.session_state.results = rows
            st.success("Resultado aceptado por el especialista.")
            st.rerun()
        if c2.button("✎ Modificar", key="mod_" + r["id"], use_container_width=True):
            rows[idx]["estado"] = new_state
            rows[idx]["obs"] = new_obs
            rows[idx]["validacion"] = "Modificado"
            st.session_state.results = rows
            st.success("Modificación registrada por el especialista.")
            st.rerun()
        if c3.button("✕ Descartar", key="des_" + r["id"], use_container_width=True):
            rows[idx]["estado"] = "Requiere revisión especializada"
            rows[idx]["obs"] = "Resultado preliminar descartado. Requiere revisión del especialista."
            rows[idx]["validacion"] = "Descartado"
            st.session_state.results = rows
            st.warning("Resultado preliminar descartado.")
            st.rerun()

    st.markdown("#### Trazabilidad del requisito seleccionado")
    st.markdown(
        f"""
        <div class="trace">
        <b>{r['id']}</b> → {r['componente']} → criterio → evidencia → fuente → resultado preliminar → validación humana
        </div>
        """,
        unsafe_allow_html=True,
    )
    nav(3, 5, "Consolidar resultados")

elif st.session_state.step == 5:
    st.header("5. Resultados y productos de salida")
    rows = st.session_state.results
    pct = progress_pct(rows)

    pendientes = sum(r.get("validacion", "Pendiente") == "Pendiente" for r in rows)
    observaciones = sum(r["estado"] != "Cumple" for r in rows)

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Avance demostrativo", f"{pct:.1f} %")
    c2.metric("Requisitos", len(rows))
    c3.metric("Requisitos con observación", observaciones)
    c4.metric("Pendientes de validación", pendientes)

    final_df = pd.DataFrame([
        {
            "ID": r["id"],
            "Componente": r["componente"],
            "Estado preliminar": r["estado"],
            "Peso (%)": r["peso"],
            "Validación humana": r.get("validacion", "Pendiente"),
            "Observación": r["obs"],
        }
        for r in rows
    ])
    st.dataframe(
        final_df,
        use_container_width=True,
        hide_index=True,
        column_config={
            "ID": st.column_config.TextColumn("ID", width="small"),
            "Componente": st.column_config.TextColumn("Componente", width="medium"),
            "Estado preliminar": st.column_config.TextColumn("Estado preliminar", width="medium"),
            "Peso (%)": st.column_config.NumberColumn("Peso (%)", format="%d"),
            "Validación humana": st.column_config.TextColumn("Validación humana", width="medium"),
            "Observación": st.column_config.TextColumn("Observación", width="large"),
        },
    )

    if pendientes > 0:
        st.warning(
            f"Validación incompleta: existen {pendientes} requisito(s) pendiente(s) de revisión humana. "
            "Los resultados pueden visualizarse y descargarse como borrador, pero la revisión no se considera finalizada."
        )
    else:
        st.success(
            "Validación humana completada para todos los requisitos del caso de prueba. "
            "Los productos generados mantienen carácter de borrador técnico."
        )

    xlsx = excel_bytes(rows)
    docx = word_bytes(rows, pct)
    a, b = st.columns(2)
    a.download_button(
        "⬇ Generar / descargar matriz Excel",
        xlsx,
        "Matriz_revision_MVP.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )
    b.download_button(
        "⬇ Generar / descargar borrador Word",
        docx,
        "Borrador_informe_MVP.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True,
    )

    st.info("Caso de uso principal recorrido de inicio a fin. Los productos de salida son borradores y permanecen sujetos a la validación técnica del especialista.")
    st.caption("Productos del MVP: matriz de revisión en Excel y borrador de informe en Word.")
    nav(4, None)

# ---------- Sidebar ----------
with st.sidebar:
    st.markdown("### Alcance del MVP")
    st.write(
        "• PDF con texto reconocible, Word y Excel\n"
        "• Evidencia y trazabilidad\n"
        "• Estados de revisión\n"
        "• Validación humana\n"
        "• Matriz Excel y borrador Word"
    )
    st.markdown("### Fuera del alcance")
    st.write(
        "• OCR\n"
        "• Análisis especializado de planos\n"
        "• Aprobación automática\n"
        "• Conformidad técnica\n"
        "• Sustitución del especialista"
    )
    st.divider()
    st.caption("Versión 2.1 · MVP final para evaluación")
